"""Read source-pptx/*.pptx and write data/deck_data.json.

    pip install python-pptx
    python3 extract_pptx.py

Only the raw slide content is pulled out here (titles, bullets, tables,
pictures). Hand-authored corrections live in data/overrides.json and are applied
in the browser, so this script can be re-run freely after replacing a .pptx.

Speaker notes are intentionally NOT extracted.
"""

import base64
import json
import os
import re

from pptx import Presentation

SRC = os.path.join(os.path.dirname(__file__), "..", "source-pptx")
FILES = [
    ("unit1", "01 Slides.pptx", "Unit 1", "Introduction"),
    ("unit2", "02.pptx", "Unit 2", "Financial Statements I"),
    ("unit3", "03.pptx", "Unit 3", "Financial Statements II"),
    ("unit4", "04.pptx", "Unit 4", "Financial Analysis"),
    ("unit5", "05.pptx", "Unit 5", "Profitability Analysis"),
]

# Placeholder shape names, Spanish (current decks) and English (in case a deck is
# re-authored in an English PowerPoint). Matched case-insensitively as prefixes.
TITLE_NAMES = ("título", "titulo", "title")
SUBTITLE_NAMES = ("subtítulo", "subtitulo", "subtitle")
BODY_NAMES = (
    "marcador de contenido", "marcador de texto", "marcador",
    "content placeholder", "text placeholder", "body",
)
NUMBER_NAMES = ("marcador de número", "marcador de numero", "slide number placeholder")
TEXTBOX_NAMES = ("textbox", "cuadro de texto", "text box")

TITLE_LAYOUTS = ("diapositiva de título", "título", "title slide", "title")

_warnings = []


def name_matches(name, options):
    n = (name or "").strip().lower()
    return any(n.startswith(o) for o in options)


def clean(t):
    if t is None:
        return ""
    t = t.replace("\x0b", "\n").replace("\x0c", "\n")
    t = re.sub(r"[ \t]+\n", "\n", t)
    return t.strip()


def para_text(p):
    return "".join(r.text for r in p.runs) or p.text


def extract_table(shape):
    tbl = shape.table
    rows = []
    for row in tbl.rows:
        cells = []
        for cell in row.cells:
            if cell.is_spanned:
                continue
            cells.append({
                "text": clean(cell.text),
                "colspan": cell.span_width if cell.is_merge_origin else 1,
                "rowspan": cell.span_height if cell.is_merge_origin else 1,
            })
        rows.append(cells)
    return rows


def img_datauri(shape):
    try:
        part = shape.image
    except Exception:
        return None
    ext = part.ext.lower()
    if ext in ("emf", "wmf"):
        return None
    mime = {
        "png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg",
        "gif": "image/gif", "bmp": "image/bmp", "tiff": "image/tiff",
    }.get(ext, "application/octet-stream")
    return "data:%s;base64,%s" % (mime, base64.b64encode(part.blob).decode("ascii"))


def add_paragraphs(shape, bullets):
    for para in shape.text_frame.paragraphs:
        t = clean(para_text(para))
        if t:
            bullets.append({"level": para.level, "text": t})


def extract_deck(key, path):
    p = Presentation(path)
    slides_out = []
    for s_i, slide in enumerate(p.slides, 1):
        layout_name = (slide.slide_layout.name or "").strip().lower()
        is_title = any(layout_name.startswith(x) for x in TITLE_LAYOUTS)

        title_text = ""
        subtitle_text = ""
        bullets, asides, tables, images = [], [], [], []
        saw_any_text = False

        for shape in slide.shapes:
            nm = shape.name

            if getattr(shape, "has_table", False):
                tables.append(extract_table(shape))

            if shape.shape_type == 13:  # PICTURE
                uri = img_datauri(shape)
                if uri:
                    images.append({"src": uri, "w": shape.width, "h": shape.height})

            if shape.has_text_frame:
                txt = clean(shape.text_frame.text)
                if not txt:
                    continue
                saw_any_text = True
                if name_matches(nm, TITLE_NAMES):
                    title_text = txt
                elif name_matches(nm, SUBTITLE_NAMES):
                    subtitle_text = txt
                elif name_matches(nm, NUMBER_NAMES):
                    pass
                elif name_matches(nm, TEXTBOX_NAMES):
                    asides.append(txt)
                elif name_matches(nm, BODY_NAMES) or shape.is_placeholder:
                    add_paragraphs(shape, bullets)
                else:
                    # unknown text shape: keep it as an aside rather than lose it
                    asides.append(txt)

        # drop speaker-cue asides like "#faithful_image"
        asides = [a for a in asides if a.strip()[:1] != "#"]

        if not (title_text or subtitle_text or bullets or asides or tables or images):
            _warnings.append("%s slide %d: nothing extracted" % (key, s_i))

        slides_out.append({
            "isTitle": is_title,
            "title": title_text,
            "subtitle": subtitle_text,
            "bullets": bullets,
            "asides": asides,
            "tables": tables,
            "images": images,
        })
    return slides_out


deck_data = {}
for key, fname, unit_no, unit_name in FILES:
    path = os.path.join(SRC, fname)
    slides = extract_deck(key, path)
    deck_data[key] = {"unit": unit_no, "name": unit_name, "slides": slides}
    print(key, "<-", fname, "->", len(slides), "slides")

out_path = os.path.join(os.path.dirname(__file__), "..", "data", "deck_data.json")
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(deck_data, f, ensure_ascii=False)

print("json size MB:", round(os.path.getsize(out_path) / 1024 / 1024, 2))
for w in _warnings:
    print("  warning:", w)
